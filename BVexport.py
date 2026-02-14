import os
import pandas as pd
import re
import urllib.parse
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.styles import Alignment

def clean(text):
    return text.strip() if text else ""

def extract_from_file(filepath):
    rows = []

    with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
        soup = BeautifulSoup(f, "html.parser")

    entries = soup.find_all("div", attrs={"data-hackid": True})

    for entry in entries:
        row = {}
        
        # -------------------------
        # Hack ID
        # -------------------------
        hack_id = entry.get("data-hackid")
        if hack_id and hack_id.isdigit():
            row["Hack ID"] = int(hack_id)
        else:
            row["Hack ID"] = hack_id

        # -------------------------
        # Title + Download
        # -------------------------
        title_tag = entry.find("a", class_="downloadLink")
        if title_tag:
            row["Title"] = clean(title_tag.text)
            row["Download Link"] = title_tag.get("href")

        # -------------------------
        # Authors
        # -------------------------
        authors = []
        author_links = []

        for a in entry.find_all("a", href=True):
            if "ByUserID" in a["href"]:
                text = clean(a.text)
                if text:
                    authors.append(text)
                    author_links.append(a["href"])

        row["Authors"] = ", ".join(authors)
        row["Author Profile Links"] = ", ".join(author_links)

        # -------------------------
        # Downloads
        # -------------------------
        downloads = entry.find("span", class_="numDownloads")
        if downloads:
            raw_downloads = clean(downloads.text)

            # Remove commas and any non-digit characters
            numeric = re.sub(r"[^\d]", "", raw_downloads)

            if numeric:
                row["Downloads"] = int(numeric)
            else:
                row["Downloads"] = numeric

        # -------------------------
        # Last Edited
        # -------------------------
        row["Last Edited"] = ""
        moderated_block = entry.find("div", class_="hackModerated")
        if moderated_block:
            # Attempt 1: directly after lastEdit <img>
            last_edit_img = moderated_block.find("img", class_="lastEdit")
            if last_edit_img and last_edit_img.next_sibling:
                date_text = last_edit_img.next_sibling.strip()
                if date_text:
                    row["Last Edited"] = date_text

            # Optional fallback: regex match anywhere in the block
            if not row["Last Edited"]:
                text = moderated_block.get_text(" ", strip=True)
                match = re.search(
                    r'\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]* \d{1,2}(?:st|nd|rd|th)?,? \d{4}\b',
                    text
                )
                if match:
                    row["Last Edited"] = match.group(0)

        # -------------------------
        # Type Line
        # -------------------------
        type_line = entry.find("span", id="typeLine")
        if type_line:
            full_type = clean(type_line.text.strip("[]"))
            parts = [p.strip() for p in full_type.split("-")]
            if len(parts) > 0:
                row["Category"] = parts[0]
            if len(parts) > 1:
                row["Action"] = parts[1]
            if len(parts) > 2:
                row["Target"] = parts[2]

        # -------------------------
        # Description
        # -------------------------
        description = entry.find("div", class_="hackInfo")
        if description:
            desc_text = clean(description.text)

            # Fix Brawl Vault's improperly escaped quotation marks
            desc_text = desc_text.replace('\\"', '"')

            row["Description"] = desc_text

        # -------------------------
        # YouTube Links
        # -------------------------
        youtube_links = []

        html_text = str(entry)

        matches = re.findall(
            r'(?:youtube\.com/(?:watch\?v=|v/|embed/)|youtu\.be/)([A-Za-z0-9_-]{11})',
            html_text
        )

        for video_id in matches:
            youtube_links.append(f"https://www.youtube.com/watch?v={video_id}")

        row["YouTube Links"] = ", ".join(set(youtube_links))

        # -------------------------
        # Images
        # -------------------------
        images = []

        # Only include links that look like images
        for a in entry.find_all("a", href=True):
            href = a["href"]
            # Check if the href ends with a common image extension
            if any(href.lower().endswith(ext) for ext in [".png", ".jpg", ".jpeg", ".gif", ".webp"]):
                images.append(href)

        row["Image URLs"] = ", ".join(images)

        # -------------------------
        # Credit (starts with "with credit to")
        # -------------------------
        credit_value = ""

        entry_text = entry.get_text(separator="\n")

        for line in entry_text.split("\n"):
            stripped = line.strip()
            if stripped.lower().startswith("with credit to"):
                credit_value = stripped
                break

        # Fix Brawl Vault's improperly escaped quotation marks
        credit_value = credit_value.replace('\\"', '"')

        row["Credit"] = credit_value
        
        # -------------------------
        # Wi-Fi Status
        # -------------------------
        entry_html = str(entry).lower()

        if "wifiyes" in entry_html:
            row["Wi-Fi"] = "Yes"
        elif "wifino" in entry_html:
            row["Wi-Fi"] = "No"
        else:
            row["Wi-Fi"] = "Unknown"
            
        # -------------------------
        # Brawl Vault Link
        # -------------------------
        brawl_link = ""
        icons_div = entry.find("div", class_="hackIcons")
        if icons_div:
            link_tag = icons_div.find("a", href=re.compile(r"Gallery/BrawlView\.php\?Number="))
            if link_tag:
                brawl_link = link_tag["href"]
        row["Brawl Vault Link"] = brawl_link
        
        rows.append(row)

    return rows


# ---------------------------------------------------
# MAIN (Uses script folder automatically)
# ---------------------------------------------------

script_folder = os.path.dirname(os.path.abspath(__file__))
all_rows = []

for file in os.listdir(script_folder):
    if file.lower().endswith(".html"):
        full_path = os.path.join(script_folder, file)
        all_rows.extend(extract_from_file(full_path))

if not all_rows:
    print("No HTML files found in this folder.")
    input("Press Enter to exit...")
    exit()

columns_order = [
    "Title", "Hack ID", "Brawl Vault Link", "Authors",
    "Author Profile Links", "Credit", "Download Link",
    "Downloads", "Last Edited",
    "Category", "Action", "Target",
    "Description", "YouTube Links", "Image URLs",
    "Wi-Fi"
]

# Remove duplicates by Hack ID (safety)
df = pd.DataFrame(all_rows, columns=columns_order)
df.drop_duplicates(subset=["Hack ID"], inplace=True)

output_path = os.path.join(
    script_folder,
    "brawl_vault_full_export.xlsx"
)

df.to_excel(output_path, index=False)

# ---------------------------------------------------
# Excel Formatting
# ---------------------------------------------------

wb = load_workbook(output_path)
ws = wb.active

# Freeze header row
ws.freeze_panes = "A2"

# Enable filters
ws.auto_filter.ref = ws.dimensions

# Left-align header row
for cell in ws[1]:
    cell.alignment = Alignment(horizontal="left")

# Auto-size columns
for column in ws.columns:
    max_length = 0
    column_letter = column[0].column_letter

    for cell in column:
        try:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        except:
            pass

    ws.column_dimensions[column_letter].width = max_length + 2

# Wrap text for Description column
for row in ws.iter_rows():
    for cell in row:
        if cell.column_letter and ws.cell(row=1, column=cell.column).value == "Description":
            cell.alignment = Alignment(wrap_text=True)

wb.save(output_path)

print("Export complete:")
print(output_path)
input("Press Enter to exit...")
